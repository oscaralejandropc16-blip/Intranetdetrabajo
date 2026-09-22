import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Paths
scratch_dir = r'C:\Users\Laptop\.gemini\antigravity-ide\brain\cbdae822-fecb-4062-9fff-bb5d5f74a6db\scratch'
brain_dir = r'C:\Users\Laptop\.gemini\antigravity-ide\brain\cbdae822-fecb-4062-9fff-bb5d5f74a6db'

prs = Presentation()

# Set to 16:9 Widescreen
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
DARK_NAVY = RGBColor(5, 11, 20)
GOLD = RGBColor(217, 160, 91)
WHITE = RGBColor(248, 250, 252)
LIGHT_GRAY = RGBColor(148, 163, 184)

def set_slide_background(slide):
    # Add a rectangle covering the entire slide to serve as background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_NAVY
    background.line.fill.background() # No line

def add_logo(slide, x, y, width):
    logo_path = os.path.join(scratch_dir, 'official_logo.png')
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, x, y, width=width)

# --- COVER SLIDE ---
slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
set_slide_background(slide)

# Cover Logo
add_logo(slide, Inches(5.16), Inches(1.5), Inches(3.0))

# Cover Title
txBox = slide.shapes.add_textbox(Inches(2), Inches(4.5), Inches(9.33), Inches(1))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "Guía Rápida de Uso — Plataforma KANT"
run.font.name = "Montserrat"
run.font.size = Pt(40)
run.font.bold = True
run.font.color.rgb = GOLD

# Cover Subtitle
txBox_sub = slide.shapes.add_textbox(Inches(2), Inches(5.5), Inches(9.33), Inches(1))
tf_sub = txBox_sub.text_frame
p_sub = tf_sub.paragraphs[0]
p_sub.alignment = PP_ALIGN.CENTER
run_sub = p_sub.add_run()
run_sub.text = "Instructivo operativo paso a paso para el registro diario de jornadas,\ncontrol de actuaciones y gestión de la intranet."
run_sub.font.name = "Montserrat"
run_sub.font.size = Pt(18)
run_sub.font.color.rgb = LIGHT_GRAY

# --- STEPS DATA ---
steps = [
    {
        "title": "1. FICHAJE (Registro de Entrada/Salida)",
        "desc": "Lo primero que deben hacer apenas comiencen su jornada laboral es ingresar a la plataforma y marcar su hora de entrada. Al finalizar su día, deben marcar su hora de salida.",
        "img": os.path.join(scratch_dir, 'shot1.png')
    },
    {
        "title": "2. LIBRO DE ACTUACIONES (Su bitácora del día)",
        "desc": "En esta pestaña van a registrar paso a paso su trabajo. Cada vez que realicen una tarea (redactar un documento, hacer una llamada, revisar un caso), deben agregar una nueva \"Actuación\". Al final del día (antes de las 6:00 PM), todo este libro se enviará automáticamente a Jefatura como su reporte de trabajo diario.",
        "img": None
    },
    {
        "title": "3. BUZÓN DE TAREAS E INSTRUCCIONES",
        "desc": "Revise constantemente este módulo. Aquí recibirá los mensajes, directrices e instrucciones procesales directas de la jefatura para su ejecución inmediata.",
        "img": os.path.join(brain_dir, 'oficial_shot4_1789152531682.png')
    },
    {
        "title": "4. CHAT DIRECTO CON JEFATURA",
        "desc": "Dentro del buzón o al abrir una tarea, tendrán acceso a un chat oficial en tiempo real. Utilícenlo para confirmar que han leído la instrucción, hacer preguntas rápidas o informar que la tarea ha sido completada exitosamente.",
        "img": os.path.join(brain_dir, '.user_uploaded', 'media_1789157448215.png')
    },
    {
        "title": "5. LIBRO DE INGRESOS",
        "desc": "En caso de que su rol lo requiera, aquí registrarán cualquier ingreso económico, cobro o facturación relacionada con los casos y trámites legales que estén gestionando para llevar un control financiero exacto.",
        "img": None
    },
    {
        "title": "6. PROGRAMACIÓN (Agenda)",
        "desc": "Este es su calendario de trabajo. Aquí deben anotar todas sus audiencias, reuniones con clientes y las fechas de vencimiento de los tribunales. Es fundamental mantener esto al día porque el sistema les enviará notificaciones y alertas para que no se les pase ninguna fecha importante.",
        "img": os.path.join(scratch_dir, 'shot2.png')
    },
    {
        "title": "7. HISTORIAL",
        "desc": "En esta pestaña podrán consultar el registro de todas las actuaciones y bitácoras que ustedes hayan enviado en días anteriores. Les servirá para buscar qué hicieron en una fecha específica o retomar tareas previas.",
        "img": None
    },
    {
        "title": "8. BIBLIOTECA / INVESTIGACIONES KANT",
        "desc": "Este módulo es nuestro repositorio jurídico. Si realizan un análisis complejo, consiguen una buena jurisprudencia o redactan un modelo de contrato excelente, deben subirlo aquí. Igualmente, pueden usar este módulo para buscar documentos de otros compañeros que les sirvan como plantilla para sus propios casos.",
        "img": None
    },
    {
        "title": "9. MÓDULO DE EXPEDIENTES",
        "desc": "Aquí encontrarán el registro de todos los casos de la firma. Su trabajo es ingresar constantemente a los expedientes que tengan asignados y registrar cada avance, novedad o cambio procesal. Esto permite que toda la firma sepa en qué estatus se encuentra un cliente sin tener que preguntar.",
        "img": os.path.join(scratch_dir, 'shot3.png')
    },
    {
        "title": "10. MÓDULO DE GASTOS",
        "desc": "Si por alguna gestión de la firma ustedes tienen que pagar taxis, copias, aranceles de notarías, etc., deben registrar ese gasto en este módulo. Es obligatorio adjuntar una foto del recibo o factura. De esta manera, Jefatura podrá revisarlo y aprobarles el reembolso del dinero.",
        "img": os.path.join(brain_dir, 'oficial_shot5_1789152553341.png')
    }
]

# --- SLIDE GENERATION LOOP ---
for step in steps:
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    set_slide_background(slide)
    
    # Top Logo
    add_logo(slide, Inches(11.3), Inches(0.4), Inches(1.5))
    
    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(10), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = step["title"]
    run.font.name = "Montserrat"
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = GOLD
    
    # Separator Line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(10.5), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD
    line.line.fill.background()
    
    has_img = step["img"] is not None and os.path.exists(step["img"])
    
    # Content Text Box
    text_width = Inches(5.5) if has_img else Inches(11.5)
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), text_width, Inches(5.0))
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    p_content = tf_content.paragraphs[0]
    p_content.space_before = Pt(10)
    p_content.line_spacing = 1.3
    run_content = p_content.add_run()
    run_content.text = step["desc"]
    run_content.font.name = "Montserrat"
    run_content.font.size = Pt(22)
    run_content.font.color.rgb = WHITE
    
    # Image
    if has_img:
        try:
            # We place the image on the right
            img_left = Inches(6.8)
            img_top = Inches(1.6)
            img_width = Inches(6.0)
            pic = slide.shapes.add_picture(step["img"], img_left, img_top, width=img_width)
            
            # Simple vertically centering based on picture height, 
            # but since height is auto-scaled, we'd need its real height. 
            # We'll just put it at top=1.6
        except Exception as e:
            pass

# --- FINAL SLIDE ---
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide)
add_logo(slide, Inches(5.16), Inches(2.5), Inches(3.0))

txBox_final = slide.shapes.add_textbox(Inches(2), Inches(4.5), Inches(9.33), Inches(1))
tf_final = txBox_final.text_frame
p_final = tf_final.paragraphs[0]
p_final.alignment = PP_ALIGN.CENTER
run_final = p_final.add_run()
run_final.text = "¡Gracias por su atención!"
run_final.font.name = "Montserrat"
run_final.font.size = Pt(44)
run_final.font.bold = True
run_final.font.color.rgb = GOLD

txBox_final_sub = slide.shapes.add_textbox(Inches(2), Inches(5.5), Inches(9.33), Inches(1))
tf_final_sub = txBox_final_sub.text_frame
p_final_sub = tf_final_sub.paragraphs[0]
p_final_sub.alignment = PP_ALIGN.CENTER
run_final_sub = p_final_sub.add_run()
run_final_sub.text = "Documento de uso interno confidencial\nRomán & Delgado | Despacho de Abogados"
run_final_sub.font.name = "Montserrat"
run_final_sub.font.size = Pt(18)
run_final_sub.font.color.rgb = LIGHT_GRAY

# Save presentation
output_path = r'c:\sistema_prueba\intranet_trabajo\Presentacion_Induccion_KANT_Premium.pptx'
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
